from pathlib import Path
import re
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).parent
OUT = ROOT / 'guides'
OUT.mkdir(exist_ok=True)
CROPS = {
    '02-apps-terminal.png': (0, 0, 438, 277),
    '04-container.png': (323, 300, 586, 74),
    '05-labs-pass.png': (323, 253, 585, 148),
    '03-vcs-fatal.png': (323, 327, 585, 158),
    '06-verdi-waveform.png': (0, 440, 529, 174),
    '07-dc-results.png': (323, 253, 585, 104),
    '08-dc-reports.png': (323, 253, 585, 236),
}

def font(style, name, size, bold=False, color='000000'):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    style.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        style.element.rPr.rFonts.attrib.pop(qn('w:'+key),None)

def hyperlink(p, label, url, code=False):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    if code:
        family = OxmlElement('w:rFonts')
        family.set(qn('w:ascii'), 'Courier New'); family.set(qn('w:hAnsi'), 'Courier New')
        props.append(family)
        size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
    run.append(props)
    value = OxmlElement('w:t'); value.text = label; run.append(value)
    link.append(run); p._p.append(link)

def inline(p, text):
    for part in re.split(r'(`[^`]+`|\[[^\]]+\]\(https?://[^)]+\)|https?://[^\s]+)', text):
        if part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Courier New'; r.font.size=Pt(9)
        elif part.startswith('http') or re.match(r'\[[^\]]+\]\(https?://', part):
            match = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', part)
            label, url = match.groups() if match else (part, part)
            hyperlink(p, label, url)
        else:
            p.add_run(part)

def make_doc(src):
    doc = Document()
    doc.core_properties.author = 'Abhinav Nandwani'
    doc.core_properties.last_modified_by = 'Abhinav Nandwani'
    doc.core_properties.title = src.read_text().splitlines()[0].removeprefix('# ')
    doc.core_properties.subject = 'CAE Synopsys terminal and GUI guide'
    doc.core_properties.keywords = 'CAE, Synopsys, Abhinav Nandwani'
    for style in doc.styles:
        if style.element.find(qn('w:pPr')) is not None:
            for border in style.element.find(qn('w:pPr')).findall(qn('w:pBdr')):
                border.getparent().remove(border)
    sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=Inches(.58); sec.bottom_margin=Inches(.58)
    sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    sec.header_distance=Inches(.25); sec.footer_distance=Inches(.25)
    font(doc.styles['Normal'], 'Arial', 10.5)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing=1.07
    for sty,size in [('Title',27),('Heading 1',17),('Heading 2',13)]:
        font(doc.styles[sty],'Arial',size,True)
        doc.styles[sty].paragraph_format.space_before=Pt(8 if sty!='Title' else 0)
        doc.styles[sty].paragraph_format.space_after=Pt(7)
    for sty in ['List Bullet','List Number']:
        font(doc.styles[sty],'Arial',10.5)
        doc.styles[sty].paragraph_format.space_after=Pt(4)
    if 'Code' not in doc.styles:
        doc.styles.add_style('Code',1)
    font(doc.styles['Code'],'Courier New',9)
    doc.styles['Code'].paragraph_format.space_after=Pt(0)
    doc.styles['Code'].paragraph_format.line_spacing=1
    doc.styles['Code'].paragraph_format.left_indent=Inches(.1)
    header=sec.header.paragraphs[0]
    header.text='ABHINAV NANDWANI  /  CAE LAB NOTES'
    for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=foot.add_run('23 September 2026   |   ');r.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    lines=src.read_text().splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='<!-- page -->':
            doc.add_page_break();i+=1;continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:],style='Title');i+=1;continue
        if line.startswith('## '):
            doc.add_paragraph(line[3:],style='Heading 1');i+=1;continue
        if line.startswith('```'):
            i+=1;code=[]
            while i<len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]);i+=1
            p=doc.add_paragraph(style='Code')
            for part in re.split(r'(https?://[^\s]+)', '\n'.join(code)):
                if part.startswith('http'):
                    hyperlink(p, part, part, code=True)
                else:
                    p.add_run(part)
            p.paragraph_format.space_after=Pt(8)
            p.paragraph_format.keep_together=True
            i+=1;continue
        if line.startswith('!['):
            m=re.match(r'!\[(.*)\]\(([^|]+)\|([\d.]+)\)',line)
            if not m:raise ValueError(line)
            caption,name,width=m.groups()
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            picture=p.add_run().add_picture(str(ROOT/'assets'/name),width=Inches(float(width)))
            # Crop the full-resolution screencap in the document itself.
            # The original PNG remains unchanged and available for inspection.
            rect=OxmlElement('a:srcRect')
            if name in CROPS:
                x,y,w,h=CROPS[name]
                iw,ih=Image.open(ROOT/'assets'/name).size
                for key,value in [('l',x/iw),('t',y/ih),('r',(iw-x-w)/iw),('b',(ih-y-h)/ih)]:
                    rect.set(key,str(round(value*100000)))
                picture.height=Inches(float(width)*h/w)
            else:
                raise ValueError(f'Missing screenshot crop: {name}')
            fill=picture._inline.graphic.graphicData.pic.blipFill
            fill.insert(1,rect)
            p=doc.add_paragraph(caption)
            p.paragraph_format.space_after=Pt(8)
            for r in p.runs:r.font.size=Pt(9);r.italic=True
            i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                vals=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',v) for v in vals):rows.append(vals)
                i+=1
            table=doc.add_table(rows=1,cols=len(rows[0]));table.autofit=False
            widths=[2.05,4.95] if len(rows[0])==2 else [7/len(rows[0])]*len(rows[0])
            for n,w in enumerate(widths):table.columns[n].width=Inches(w)
            for ri,row in enumerate(rows):
                cells=table.rows[0].cells if ri==0 else table.add_row().cells
                trPr=table.rows[ri]._tr.get_or_add_trPr()
                noSplit=OxmlElement('w:cantSplit');trPr.append(noSplit)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
                for ci,val in enumerate(row):
                    c=cells[ci];c.width=Inches(widths[ci])
                    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                    inline(p,val)
                    for r in p.runs:r.font.size=Pt(9.4);r.bold=ri==0
                    tcPr=c._tc.get_or_add_tcPr()
                    marg=OxmlElement('w:tcMar')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');marg.append(e)
                    tcPr.append(marg)
                    borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D5D5D5');borders.append(e)
                    tcPr.append(borders)
                    if ri==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');tcPr.append(shade)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            continue
        m=re.match(r'\d+\. (.*)',line)
        p=doc.add_paragraph(style='Normal')
        if m:
            p.paragraph_format.left_indent=Inches(.2)
            p.paragraph_format.first_line_indent=Inches(-.2)
            p.paragraph_format.space_after=Pt(4)
        inline(p,line)
        if line.startswith('By Abhinav Nandwani'):
            for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('505050')
        i+=1
    target=OUT/(src.stem+'.docx')
    doc.save(target)
    print(target)

if __name__=='__main__':
    raise SystemExit('Run python3 build_guides.py to build all three guides.')
